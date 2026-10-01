"""Generate testdata/products.xlsx. Run once:  python -m utils.create_test_excel

Edit the workbook directly afterwards to add/remove scenarios - no code change needed.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "testdata" / "products.xlsx"

FONT = "Arial"
HEADERS = ["TestCase ID", "Run", "Search Term", "Product Name", "Update Quantity"]
ROWS = [
    ["TC01_MacBook", "Y", "MacBook", "MacBook", 3],
    ["TC02_iPhone", "Y", "iPhone", "iPhone", 2],
    ["TC03_iPod_Nano", "Y", "iPod", "iPod Nano", 4],
    ["TC04_HTC_Touch_HD", "N", "HTC", "HTC Touch HD", 2],  # switched off with Run = N
]

LEGEND = [
    ("Products sheet - how to use", None),
    ("TestCase ID", "Unique name; shown in the pytest report."),
    ("Run", "Y = execute this row, N = skip it."),
    ("Search Term", "Text typed in the store search box."),
    ("Product Name", "Exact product title to add to the cart (must appear in the results)."),
    ("Update Quantity", "Quantity to set on the cart page (whole number, 1 or more)."),
    ("Note", "Only add products that have no required options (e.g. not 'Apple Cinema 30\"')."),
]


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)

    head_fill = PatternFill("solid", start_color="1F3864")
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    for col, width in zip("ABCDE", (22, 8, 16, 22, 18)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B2:B200")

    help_ws = wb.create_sheet("Instructions")
    for r, (a, b) in enumerate(LEGEND, start=1):
        help_ws.cell(row=r, column=1, value=a).font = Font(name=FONT, bold=True, size=12 if r == 1 else 10)
        if b:
            help_ws.cell(row=r, column=2, value=b).font = Font(name=FONT)
    help_ws.column_dimensions["A"].width = 30
    help_ws.column_dimensions["B"].width = 80

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    build()
