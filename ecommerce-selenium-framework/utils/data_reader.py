"""Read framework config and test data from JSON and Excel files."""
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CONFIG_FILE = ROOT / "config" / "config.json"
JSON_DATA_FILE = ROOT / "testdata" / "test_data.json"
EXCEL_DATA_FILE = ROOT / "testdata" / "products.xlsx"


def read_json(path: Path) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def get_config() -> dict:
    return read_json(CONFIG_FILE)


def get_json_test_data() -> dict:
    return read_json(JSON_DATA_FILE)


def read_excel_rows(path: Path = EXCEL_DATA_FILE, sheet: str = "Products") -> list[dict]:
    """Return every non-empty row of `sheet` as a dict keyed by the header row.

    Rows whose `Run` column is 'N' (case-insensitive) are skipped, so testers can
    switch scenarios on/off from Excel without touching code.
    """
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() for h in rows[0]]
    data = []
    for row in rows[1:]:
        if all(cell is None for cell in row):
            continue
        record = dict(zip(headers, row))
        if str(record.get("Run", "Y")).strip().upper() == "N":
            continue
        data.append(record)
    return data
